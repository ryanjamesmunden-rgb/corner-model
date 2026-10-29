"""Turn the form and fixtures exports into one fill-in EV workbook.

WHY THIS AND NOT board_xlsx.py. That one takes the BOARD — rows that already qualified
under the board's own rules — and prices them. These two tables are the opposite shape on
purpose: every team and every upcoming side, including the ones no angle picked, because a
sheet is read by sorting and filtering it and you cannot filter rows that were dropped
before you saw them.

THE EV COLUMN IS A FORMULA, NEVER A NUMBER, for the reason board_xlsx.py gives: a
bookmaker's price is the one thing the model cannot know, so the sheet arrives with that
column empty and the edge appears as it is typed in. A sheet with EV pre-computed for the
few rows that happen to be priced and blanks elsewhere is a dead document.

AND IT IS GUARDED WITH IFERROR AND AN EMPTY-STRING TEST. A blank price would otherwise
evaluate to -100% on every unpriced row, which reads as a strong call AGAINST the bet
rather than as an absence of information.

Reads the two CSVs that /api/tools/sheet.csv returns; writes one .xlsx. No network.

Run: python tools/sheet_xlsx.py form.csv fixtures.csv out.xlsx
"""
import csv
import sys
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")      # the columns you type in
THIN_FILL = PatternFill("solid", fgColor="FCE4D6")       # a thin sample, flagged
AWAY_FILL = PatternFill("solid", fgColor="EDEDED")       # an away row, flagged
BLACK = Font(name=FONT, size=10)
BLUE = Font(name=FONT, color="0000FF", size=10)
BOLD = Font(name=FONT, bold=True, size=10)
GREY = Font(name=FONT, size=9, italic=True, color="808080")
WARN = Font(name=FONT, size=9, italic=True, color="C00000")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LINES = (5, 6, 7, 8)

# The haircut the record actually supports, applied as a SHEET-LEVEL input rather than
# baked into a column. The 1,919-row walk-forward replay of the mismatch board puts the
# residual at -4.7pt on home rows and -7.1 on away; the 82-row frozen record says +0.3.
# Those disagree, and the disagreement is worth about seven points of probability — which
# is the difference between a handful of bets and most of the card. So it is a cell: put
# 0 in it to price off the model as published, or the replay's figure to be conservative.
DEFAULT_SHADE_HOME = 4.7
DEFAULT_SHADE_AWAY = 7.1
EDGE_TARGET = 5.0        # the +5% the overs sheet has always asked for


def read_csv(path):
    if not path:
        return []
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def num(v):
    if v in (None, "", "None"):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f == int(f) else f


def head(ws, labels, widths):
    for i, (label, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30


def legend(ws, row, lines):
    for i, text in enumerate(lines):
        c = ws.cell(row=row + i, column=1, value=text)
        c.font = WARN if text.startswith("!") else GREY


# ---------------------------------------------------------------------------


def form_tab(wb, rows, window):
    """Recent corner form, every team. The question: who is winning corners right now."""
    ws = wb.create_sheet("Form — last 5")
    cols = ["Team", "League", "Games", f"Last {window} avg", f"Last {window}",
            "Low", "High", "Last 10 avg", "Conceded 10", "Home avg", "Away avg"]
    widths = [24, 20, 7, 10, 14, 6, 6, 10, 11, 10, 10]
    for line in LINES:
        cols += [f"{line}+ in {window}", f"{line}+ in 10", f"{line}+ run"]
        widths += [9, 9, 8]
    head(ws, cols, widths)
    ws.freeze_panes = "C2"

    for r, row in enumerate(rows, start=2):
        vals = [row.get("team"), row.get("league"), num(row.get("games")),
                num(row.get(f"last{window}_avg")), row.get(f"last{window}"),
                num(row.get(f"last{window}_low")), num(row.get(f"last{window}_high")),
                num(row.get("last10_avg")), num(row.get("conceded_last10")),
                num(row.get("home_avg")), num(row.get("away_avg"))]
        for line in LINES:
            vals += [num(row.get(f"hit{line}_of{window}")),
                     num(row.get(f"hit{line}_of10")), num(row.get(f"run{line}"))]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.border = BLACK, BORDER
            if i in (4, 8, 9, 10, 11):
                c.number_format = "0.00"
        # A SHORT RECORD IS FLAGGED, NOT DROPPED. A five-game average off six games on
        # file is a real average over a thin season, and the reader should be able to see
        # which it is without opening another sheet.
        if (num(row.get("games")) or 0) < 10:
            for i in range(1, len(vals) + 1):
                ws.cell(row=r, column=i).fill = THIN_FILL

    legend(ws, len(rows) + 3, [
        f"Corners WON, over each side's last {window} real games — newest first in the "
        f"'Last {window}' column. Backward-looking: nothing here is a tip.",
        f"'{LINES[0]}+ in {window}' is a HIT COUNT and '{LINES[0]}+ run' is the CURRENT "
        "run. They are different claims: 9 of 10 tolerates a miss, 9 in a row does not, "
        "and a side can be 9/10 with a run of 2.",
        "Shaded rows have under 10 games on file — the average is real, the sample is thin.",
        "Real matches only. Padded games are excluded, so every figure here can be "
        "checked against a scoreline.",
    ])
    return ws


def fixtures_tab(wb, rows, window):
    """Every upcoming side with the 5+ to 8+ ladder, and the fill-in half."""
    ws = wb.create_sheet("Fixtures — next 7")

    # The two shade cells every row reads. See DEFAULT_SHADE_* for why this is an input.
    ws["A1"] = "Shade home (pts)"
    ws["B1"] = DEFAULT_SHADE_HOME
    ws["C1"] = "Shade away (pts)"
    ws["D1"] = DEFAULT_SHADE_AWAY
    ws["E1"] = "Target edge %"
    ws["F1"] = EDGE_TARGET
    for cell in ("A1", "C1", "E1"):
        ws[cell].font = BOLD
    for cell in ("B1", "D1", "F1"):
        ws[cell].fill, ws[cell].font, ws[cell].number_format = INPUT_FILL, BLUE, "0.0"
    ws["G1"] = ("Model probability minus the shade for that venue. 0 = price off the "
                "model as published.")
    ws["G1"].font = GREY

    cols = ["Kick-off", "League", "Team", "", "Opponent", "Lambda",
            f"Venue last {window}", "Games", "Line", "Model %", "Shaded %",
            "Fair", "Target", "Book", "EV", "Verdict"]
    widths = [16, 18, 22, 6, 22, 8, 14, 7, 6, 9, 10, 8, 8, 9, 9, 16]
    HEAD_ROW = 3
    for i, (label, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=HEAD_ROW, column=i, value=label)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[HEAD_ROW].height = 30
    ws.freeze_panes = f"A{HEAD_ROW + 1}"

    r = HEAD_ROW + 1
    for row in rows:
        side = (row.get("side") or "").lower()
        is_away = side == "away"
        for line in LINES:
            p = num(row.get(f"p{line}"))
            if p is None:
                continue
            base = [
                row.get("kickoff"), row.get("league"), row.get("team"),
                "away" if is_away else "home", row.get("opponent"),
                num(row.get("lambda")), num(row.get(f"venue_last{window}_avg")),
                num(row.get("games")), f"{line}+",
            ]
            for i, v in enumerate(base, start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.font, c.border = BLACK, BORDER
                if i in (6, 7):
                    c.number_format = "0.00"
            ws.cell(row=r, column=10, value=p / 100).number_format = "0.0%"
            # Shaded probability reads the venue's own cell, so changing one number at the
            # top re-prices the whole sheet rather than needing a re-export.
            shade = "$B$1" if not is_away else "$D$1"
            ws[f"K{r}"] = f"=MAX(0,J{r}-{shade}/100)"
            ws[f"K{r}"].number_format = "0.0%"
            ws.cell(row=r, column=12, value=num(row.get(f"fair{line}"))).number_format = "0.00"
            # Target is off the SHADED probability, not the published one: a price that
            # only clears the model's own number is not the bar this sheet asks for.
            ws[f"M{r}"] = f'=IFERROR(IF(K{r}<=0,"",1/K{r}*(1+$F$1/100)),"")'
            ws[f"M{r}"].number_format = "0.00"
            b = ws.cell(row=r, column=14)
            b.fill, b.font, b.number_format, b.border = INPUT_FILL, BLUE, "0.00", BORDER
            ws[f"O{r}"] = f'=IFERROR(IF(N{r}="","",N{r}*K{r}-1),"")'
            ws[f"O{r}"].number_format = "0.0%;[Red]-0.0%;-"
            ws[f"P{r}"] = (
                f'=IF(O{r}="","awaiting price",'
                f'IF(N{r}<M{r},"under target",'
                f'IF(O{r}>=$F$1/100,"TAKE","thin")))')
            ws[f"P{r}"].font = BLACK
            for i in range(10, 17):
                ws.cell(row=r, column=i).border = BORDER
            if is_away:
                for i in range(1, 17):
                    ws.cell(row=r, column=i).fill = (
                        INPUT_FILL if i == 14 else AWAY_FILL)
            r += 1

    legend(ws, r + 2, [
        "One row per side per line. Type a price in the yellow Book column and EV "
        "appears; nothing else needs touching.",
        "Model % comes from nb_ge, the same function the site prices with. Shaded % "
        "subtracts the venue's cell at the top — edit those two numbers to re-price "
        "every row at once.",
        "Target is the price that clears the shaded probability by the target edge. "
        "'under target' means the book price is short of it.",
        "! Away rows are shaded grey. In the frozen quick-scan record they are 4 of 8 "
        "settled, and the 1,919-row replay agrees on the direction — the weakest arm "
        "on the board. Eight games is not enough to drop them, so they are marked.",
        "! 8+ needs lambda around 8.5 before it beats even money, and the record holds "
        "exactly one settled 8+ claim, which lost. That line has no evidence either way.",
        "No book odds are supplied. There is no price feed, so EV depends only on what "
        "you type.",
    ])
    return ws


def cover(wb, form_rows, fixture_rows, window, days):
    ws = wb.create_sheet("Read me", 0)
    ws.column_dimensions["A"].width = 112
    built = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    lines = [
        ("Corner EV sheet", BOLD),
        (f"Built {built}. Form over the last {window} games; fixtures over the next "
         f"{days} days.", BLACK),
        ("", BLACK),
        (f"Form — last {window}:  {len(form_rows)} teams, every league. Corners won "
         "recently, with a hit count AND a current run at each line from 5+ to 8+.",
         BLACK),
        (f"Fixtures — next {days}:  {len(fixture_rows)} sides, each priced at 5+, 6+, "
         "7+ and 8+. Type a book price, read the EV.", BLACK),
        ("", BLACK),
        ("What this sheet does NOT contain", BOLD),
        ("Book odds. There is no price feed, so every EV here depends only on the price "
         "you type in.", BLACK),
        ("Any claim about a team's history beyond its stored record — roughly twenty "
         "real matches. Nothing here reaches back further.", BLACK),
        ("", BLACK),
        ("Two things the record says before you stake", BOLD),
        ("! Away rows are the weak arm: 4 of 8 settled in the frozen quick-scan record, "
         "and the 1,919-row replay points the same way. They are marked, not removed.",
         WARN),
        ("! The shade is a judgement, not a result. The replay says -4.7pt home and "
         "-7.1 away; the 82-row frozen record says +0.3. Both are on the sheet as "
         "editable cells because they disagree.", WARN),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def main():
    if len(sys.argv) < 4:
        print(__doc__)
        return 2
    form_path, fix_path, out = sys.argv[1], sys.argv[2], sys.argv[3]
    window = int(sys.argv[4]) if len(sys.argv) > 4 else 5
    days = int(sys.argv[5]) if len(sys.argv) > 5 else 7

    form_rows = read_csv(form_path)
    fixture_rows = read_csv(fix_path)
    if not form_rows and not fixture_rows:
        print("nothing to write — both exports were empty")
        return 1

    wb = Workbook()
    wb.remove(wb.active)
    cover(wb, form_rows, fixture_rows, window, days)
    if form_rows:
        form_tab(wb, form_rows, window)
    if fixture_rows:
        fixtures_tab(wb, fixture_rows, window)
    wb.save(out)
    print(f"wrote {out}: {len(form_rows)} form rows, {len(fixture_rows)} fixture sides")
    return 0


if __name__ == "__main__":
    sys.exit(main())
